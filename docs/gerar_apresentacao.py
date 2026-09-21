import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from front.componentes.grafico_confusao import montar_grafico_matriz_confusao
from front.dados.consultas import listar_execucoes, obter_gabarito, obter_metricas, obter_parametros
from front.dados.consultas_ml import obter_ultimo_resultado_ml
from front.dados.derivadas import calcular_comparacao_binaria_ml, calcular_indicadores_por_algoritmo

CAMINHO_SAIDA = Path(__file__).resolve().parent / "apresentacao_wearguard.pptx"
CAMINHO_IMAGEM_MATRIZ = Path(__file__).resolve().parent / "_matriz_confusao_temp.png"

COR_DESTAQUE = RGBColor(0x25, 0x63, 0xEB)
COR_LIMIAR_SIMPLES = RGBColor(0xF9, 0x73, 0x16)
COR_PERSISTENCIA = RGBColor(0x16, 0xA3, 0x4A)
COR_TEXTO = RGBColor(0x1E, 0x29, 0x3B)
COR_TEXTO_SUAVE = RGBColor(0x64, 0x74, 0x8B)
COR_FUNDO = RGBColor(0xF6, 0xF8, 0xFB)

LARGURA_SLIDE = Inches(13.333)
ALTURA_SLIDE = Inches(7.5)


def obter_dados_regras() -> dict | None:
    """Busca os parâmetros, o gabarito e os indicadores da execução simulada mais recente."""
    execucoes = listar_execucoes()
    if not execucoes:
        return None
    execucao_id = execucoes[0][0]
    metricas = obter_metricas(execucao_id)
    return {"parametros": obter_parametros(execucao_id), "gabarito": obter_gabarito(execucao_id), "indicadores": calcular_indicadores_por_algoritmo(metricas)}


def criar_apresentacao() -> Presentation:
    """Cria uma apresentação em branco no formato widescreen (16:9)."""
    prs = Presentation()
    prs.slide_width = LARGURA_SLIDE
    prs.slide_height = ALTURA_SLIDE
    return prs


def pintar_fundo(slide, cor: RGBColor) -> None:
    """Preenche o fundo do slide com uma cor sólida."""
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = cor


def adicionar_slide_branco(prs: Presentation):
    """Adiciona um slide em branco (sem placeholders) para montagem manual."""
    layout_em_branco = prs.slide_layouts[6]
    return prs.slides.add_slide(layout_em_branco)


def adicionar_titulo_slide(slide, texto: str, cor: RGBColor = COR_TEXTO, tamanho: int = 32) -> None:
    """Adiciona uma caixa de título no topo do slide."""
    caixa = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), LARGURA_SLIDE - Inches(1.2), Inches(1))
    paragrafo = caixa.text_frame.paragraphs[0]
    execucao = paragrafo.add_run()
    execucao.text = texto
    execucao.font.size = Pt(tamanho)
    execucao.font.bold = True
    execucao.font.color.rgb = cor


def adicionar_barra_destaque(slide, cor: RGBColor) -> None:
    """Adiciona uma barra colorida fina no topo do slide, para identidade visual."""
    barra = slide.shapes.add_shape(1, Inches(0), Inches(0), LARGURA_SLIDE, Inches(0.15))
    barra.fill.solid()
    barra.fill.fore_color.rgb = cor
    barra.line.fill.background()


def adicionar_slide_capa(prs: Presentation, titulo: str, subtitulo: str) -> None:
    """Monta o slide de capa da apresentação."""
    slide = adicionar_slide_branco(prs)
    pintar_fundo(slide, COR_FUNDO)
    adicionar_barra_destaque(slide, COR_DESTAQUE)

    caixa_titulo = slide.shapes.add_textbox(Inches(0.8), Inches(2.8), LARGURA_SLIDE - Inches(1.6), Inches(1.5))
    execucao_titulo = caixa_titulo.text_frame.paragraphs[0].add_run()
    execucao_titulo.text = titulo
    execucao_titulo.font.size = Pt(44)
    execucao_titulo.font.bold = True
    execucao_titulo.font.color.rgb = COR_TEXTO

    caixa_subtitulo = slide.shapes.add_textbox(Inches(0.8), Inches(4.2), LARGURA_SLIDE - Inches(1.6), Inches(1))
    execucao_subtitulo = caixa_subtitulo.text_frame.paragraphs[0].add_run()
    execucao_subtitulo.text = subtitulo
    execucao_subtitulo.font.size = Pt(20)
    execucao_subtitulo.font.italic = True
    execucao_subtitulo.font.color.rgb = COR_TEXTO_SUAVE


def adicionar_slide_bullets(prs: Presentation, titulo: str, bullets: list[str], cor_barra: RGBColor = COR_DESTAQUE) -> None:
    """Adiciona um slide com título e uma lista de bullets."""
    slide = adicionar_slide_branco(prs)
    pintar_fundo(slide, RGBColor(0xFF, 0xFF, 0xFF))
    adicionar_barra_destaque(slide, cor_barra)
    adicionar_titulo_slide(slide, titulo)

    caixa = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), LARGURA_SLIDE - Inches(1.6), ALTURA_SLIDE - Inches(2.2))
    quadro_texto = caixa.text_frame
    quadro_texto.word_wrap = True

    for indice, bullet in enumerate(bullets):
        paragrafo = quadro_texto.paragraphs[0] if indice == 0 else quadro_texto.add_paragraph()
        execucao = paragrafo.add_run()
        execucao.text = f"•  {bullet}"
        execucao.font.size = Pt(19)
        execucao.font.color.rgb = COR_TEXTO
        paragrafo.space_after = Pt(14)


def adicionar_slide_tabela(prs: Presentation, titulo: str, colunas: list[str], linhas: list[list[str]], cor_barra: RGBColor = COR_DESTAQUE) -> None:
    """Adiciona um slide com título e uma tabela de dados."""
    slide = adicionar_slide_branco(prs)
    pintar_fundo(slide, RGBColor(0xFF, 0xFF, 0xFF))
    adicionar_barra_destaque(slide, cor_barra)
    adicionar_titulo_slide(slide, titulo)

    numero_linhas = len(linhas) + 1
    numero_colunas = len(colunas)
    tabela_shape = slide.shapes.add_table(numero_linhas, numero_colunas, Inches(0.8), Inches(1.7), LARGURA_SLIDE - Inches(1.6), Inches(0.5) * numero_linhas)
    tabela = tabela_shape.table

    for indice_coluna, nome_coluna in enumerate(colunas):
        celula = tabela.cell(0, indice_coluna)
        celula.text = nome_coluna
        celula.text_frame.paragraphs[0].runs[0].font.bold = True
        celula.text_frame.paragraphs[0].runs[0].font.size = Pt(15)

    for indice_linha, linha in enumerate(linhas, start = 1):
        for indice_coluna, valor in enumerate(linha):
            celula = tabela.cell(indice_linha, indice_coluna)
            celula.text = str(valor)
            celula.text_frame.paragraphs[0].runs[0].font.size = Pt(14)


def adicionar_slide_imagem(prs: Presentation, titulo: str, caminho_imagem: Path, cor_barra: RGBColor = COR_DESTAQUE) -> None:
    """Adiciona um slide com título e uma imagem centralizada."""
    slide = adicionar_slide_branco(prs)
    pintar_fundo(slide, RGBColor(0xFF, 0xFF, 0xFF))
    adicionar_barra_destaque(slide, cor_barra)
    adicionar_titulo_slide(slide, titulo)
    slide.shapes.add_picture(str(caminho_imagem), Inches(3.2), Inches(1.5), height = Inches(5.5))


def montar_slides_capa_e_visao_geral(prs: Presentation) -> None:
    adicionar_slide_capa(prs, "WearGuard", "Uso de wearables e sensores no monitoramento remoto de pacientes")
    adicionar_slide_bullets(prs, "Visão geral do projeto", [
        "Frente 1 (escopo original): simulador de sinais vitais + dois algoritmos de detecção de anomalias",
        "Frente 2 (extensão): classificador de arritmias com machine learning, treinado em ECG real",
        "As duas frentes usam dados e resolvem problemas diferentes",
        "Frente 1 decide QUANDO alertar; Frente 2 decide QUE TIPO de batimento está ocorrendo",
    ])


def montar_slides_frente_regras(prs: Presentation, dados_regras: dict | None) -> None:
    adicionar_slide_bullets(prs, "Frente 1 — Por que dados simulados?", [
        "Dados de saúde são protegidos pela LGPD",
        "Dados reais exigiriam aprovação em Comitê de Ética / Plataforma Brasil",
        "Prazo do TCC incompatível com esse processo",
        "Simulação permite controlar exatamente onde cada evento acontece (o gabarito)",
    ], cor_barra = COR_LIMIAR_SIMPLES)

    adicionar_slide_bullets(prs, "Frente 1 — Como o simulador funciona", [
        "Gera um sinal-base saudável para frequência cardíaca, SpO2 e temperatura",
        "Aplica ruído gaussiano pequeno (imprecisão natural de sensores)",
        "Insere eventos anômalos reais em instantes sorteados (taquicardia, bradicardia, hipóxia, febre)",
        "Registra o gabarito: onde cada evento começa e termina",
    ], cor_barra = COR_LIMIAR_SIMPLES)

    adicionar_slide_bullets(prs, "Frente 1 — Os dois algoritmos", [
        "Limiar simples: alerta na hora, em qualquer leitura fora da faixa normal",
        "Persistência: só alerta após 5 leituras seguidas fora da faixa",
        "Hipótese: persistência reduz falsos positivos, ao custo de um tempo de detecção maior",
    ], cor_barra = COR_LIMIAR_SIMPLES)

    if dados_regras is not None and dados_regras["indicadores"]:
        colunas = ["Algoritmo", "Verdadeiros positivos", "Falsos positivos", "Falsos negativos", "Tempo médio (s)"]
        linhas = []
        for algoritmo, valores in dados_regras["indicadores"].items():
            nome = "Limiar simples" if algoritmo == "limiar_simples" else "Persistência"
            linhas.append([nome, valores["verdadeiros_positivos"], valores["falsos_positivos"], valores["falsos_negativos"], valores["tempo_medio_deteccao"]])
        adicionar_slide_tabela(prs, "Frente 1 — Resultados obtidos", colunas, linhas, cor_barra = COR_LIMIAR_SIMPLES)


def montar_slides_frente_ml(prs: Presentation, resultado_ml: dict | None) -> None:
    adicionar_slide_bullets(prs, "Frente 2 — Machine Learning com ECG real", [
        "Adicionada depois, fora do escopo original do projeto de pesquisa",
        "Objetivo: comparar regras fixas com uma abordagem que aprende de dados reais",
        "Usa só o sinal elétrico do coração (ECG) — não usa SpO2 nem temperatura",
        "Base: MIT-BIH Arrhythmia Database (PhysioNet), 48 registros reais, 47 pacientes",
    ], cor_barra = COR_PERSISTENCIA)

    adicionar_slide_bullets(prs, "Frente 2 — Metodologia", [
        "Cada batimento: janela de sinal ao redor do pico R + intervalos RR (tempo até o batimento anterior/seguinte)",
        "5 classes da norma AAMI EC57: normal, supraventricular, ventricular, fusão, desconhecido",
        "Treino e teste separados POR PACIENTE (22 + 22), técnica de de Chazal et al. (2004)",
        "Evita que o modelo memorize características de um paciente específico",
    ], cor_barra = COR_PERSISTENCIA)

    adicionar_slide_bullets(prs, "Frente 2 — Algoritmo: Random Forest", [
        "Floresta de árvores de decisão, com balanceamento de classes",
        "Robusto, rápido de treinar, sem necessidade de GPU",
        "Bem documentado na literatura de classificação de ECG",
        "Mais fácil de explicar e defender que uma rede neural profunda",
    ], cor_barra = COR_PERSISTENCIA)

    if resultado_ml is not None:
        colunas = ["Classe", "Precisão", "Recall", "F1", "Suporte no teste"]
        nomes = {"normal": "Normal", "supraventricular": "Supraventricular", "ventricular": "Ventricular", "fusao": "Fusão", "desconhecido": "Desconhecido"}
        linhas = []
        for _, linha in resultado_ml["metricas_por_classe"].iterrows():
            linhas.append([nomes.get(linha["classe"], linha["classe"]), f"{linha['precisao']:.3f}", f"{linha['recall']:.3f}", f"{linha['f1']:.3f}", int(linha["suporte"])])
        adicionar_slide_tabela(prs, f"Frente 2 — Resultados (acurácia {resultado_ml['acuracia'] * 100:.1f}%, macro-F1 {resultado_ml['macro_f1']:.3f})", colunas, linhas, cor_barra = COR_PERSISTENCIA)

        figura = montar_grafico_matriz_confusao(resultado_ml["matriz_confusao"])
        figura.savefig(CAMINHO_IMAGEM_MATRIZ, dpi = 150, bbox_inches = "tight")
        adicionar_slide_imagem(prs, "Frente 2 — Matriz de confusão", CAMINHO_IMAGEM_MATRIZ, cor_barra = COR_PERSISTENCIA)


def montar_slide_comparacao(prs: Presentation, dados_regras: dict | None, resultado_ml: dict | None) -> None:
    adicionar_slide_bullets(prs, "Comparação entre as duas frentes", [
        "Problemas e dados diferentes: não dá para comparar acurácia com falsos positivos direto",
        "Solução: agrupar as 4 classes de arritmia em \"anômalo\" vs. \"normal\"",
        "Assim as duas frentes respondem a mesma pergunta binária: falso alarme e detecção",
    ])

    if dados_regras is not None and dados_regras["indicadores"] and resultado_ml is not None:
        comparacao_ml = calcular_comparacao_binaria_ml(resultado_ml["matriz_confusao"])
        colunas = ["Abordagem", "Taxa de falso alarme", "Taxa de detecção"]
        linhas = []
        for algoritmo, valores in dados_regras["indicadores"].items():
            nome = "Limiar simples" if algoritmo == "limiar_simples" else "Persistência"
            linhas.append([nome, f"{valores['taxa_falsos_positivos'] * 100:.1f}%", f"{valores['taxa_deteccao'] * 100:.1f}%"])
        linhas.append(["Machine Learning (binarizado)", f"{comparacao_ml['taxa_falsos_positivos'] * 100:.1f}%", f"{comparacao_ml['taxa_deteccao'] * 100:.1f}%"])
        adicionar_slide_tabela(prs, "Comparação direta: quem se sai melhor", colunas, linhas)


def montar_slide_conclusao(prs: Presentation) -> None:
    adicionar_slide_bullets(prs, "Pontos fortes do trabalho", [
        "Base de dados padrão-ouro da literatura (MIT-BIH), com split correto por paciente",
        "Reprodutibilidade: semente aleatória fixa, código versionado, dataset público",
        "Duas abordagens comparadas de forma controlada e justa",
        "Dashboard interativo para visualizar sinais, alertas e métricas",
        "Discussão honesta das limitações, com trabalhos futuros apontados na literatura",
    ])


def gerar_apresentacao() -> None:
    """Gera a apresentação completa (PPTX) com os dados reais das duas frentes do projeto."""
    dados_regras = obter_dados_regras()
    resultado_ml = obter_ultimo_resultado_ml()

    prs = criar_apresentacao()
    montar_slides_capa_e_visao_geral(prs)
    montar_slides_frente_regras(prs, dados_regras)
    montar_slides_frente_ml(prs, resultado_ml)
    montar_slide_comparacao(prs, dados_regras, resultado_ml)
    montar_slide_conclusao(prs)

    CAMINHO_SAIDA.parent.mkdir(parents = True, exist_ok = True)
    prs.save(str(CAMINHO_SAIDA))

    if CAMINHO_IMAGEM_MATRIZ.exists():
        CAMINHO_IMAGEM_MATRIZ.unlink()

    print(f"Apresentacao gerada em {CAMINHO_SAIDA}")


if __name__ == "__main__":
    gerar_apresentacao()
